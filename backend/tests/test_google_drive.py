import pytest
from googleapiclient.errors import HttpError

from app.services import google_drive
from app.services.google_drive import DriveError


class FakeResp:
    def __init__(self, status):
        self.status = status
        self.reason = "x"


class FakeRequest:
    def __init__(self, result=None, error=None):
        self.result, self.error = result, error

    def execute(self):
        if self.error:
            raise self.error
        return self.result


class FakeFiles:
    def __init__(self, pages, exports=None, error=None):
        self.pages, self.exports, self.error = pages, exports or {}, error
        self.calls = 0

    def list(self, **kwargs):
        assert "mimeType='application/vnd.google-apps.document'" in kwargs["q"]
        if self.error:
            return FakeRequest(error=self.error)
        page = self.pages[self.calls]
        self.calls += 1
        return FakeRequest(page)

    def export(self, fileId, mimeType):
        r = self.exports[fileId]
        return FakeRequest(error=r) if isinstance(r, Exception) else FakeRequest(r)


class FakeService:
    def __init__(self, files):
        self._files = files

    def files(self):
        return self._files


def test_pagination_and_fields():
    pages = [
        {"files": [{"id": "a", "name": "A", "modifiedTime": "t1", "webViewLink": "http://a"}], "nextPageToken": "n"},
        {"files": [{"id": "b", "name": "B", "modifiedTime": "t2"}]},
    ]
    docs = google_drive.list_documents(FakeService(FakeFiles(pages)), "folder")
    assert [d["id"] for d in docs] == ["a", "b"]
    assert docs[0]["url"] == "http://a"
    assert docs[1]["url"] == "https://docs.google.com/document/d/b/edit"


def test_missing_folder_id(monkeypatch):
    monkeypatch.setattr(google_drive.settings, "google_drive_folder_id", "")
    with pytest.raises(DriveError, match="FOLDER_ID"):
        google_drive.list_documents(FakeService(FakeFiles([])))


def test_permission_error_message_has_no_internals():
    err = HttpError(FakeResp(403), b"{}")
    with pytest.raises(DriveError, match="Permission denied"):
        google_drive.list_documents(FakeService(FakeFiles([], error=err)), "folder")


def test_export_decodes_and_handles_errors():
    svc = FakeService(FakeFiles([], exports={"ok": b"hello", "bad": HttpError(FakeResp(404), b"{}")}))
    assert google_drive.read_document_text("ok", svc) == "hello"
    with pytest.raises(DriveError, match="not found"):
        google_drive.read_document_text("bad", svc)


def test_docx_extraction_includes_paragraphs_and_tables():
    import io

    from docx import Document

    d = Document()
    d.add_paragraph("Invoice requirements overview")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Field"
    t.rows[0].cells[1].text = "Amount"
    buf = io.BytesIO()
    d.save(buf)
    text = google_drive.extract_docx_text(buf.getvalue())
    assert "Invoice requirements overview" in text and "Field | Amount" in text


def test_docx_corrupt_file_gives_friendly_error():
    with pytest.raises(DriveError, match="Word file"):
        google_drive.extract_docx_text(b"not a zip")


def test_missing_token(monkeypatch, tmp_path):
    monkeypatch.setattr(google_drive.settings, "google_token_file", str(tmp_path / "nope.json"))
    with pytest.raises(DriveError, match="not authorized"):
        google_drive.build_service()
